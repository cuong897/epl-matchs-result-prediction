"""Author a Vietnamese academic report from frozen results using python-docx.

Run with the bundled Documents Python runtime. Never fits or evaluates a model.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from delivery_data import load_delivery,FAMILY

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports/delivery'


def main():
    data=load_delivery(ROOT);r=data['result'];metrics=r['metrics']
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(2.1);section.bottom_margin=Cm(1.9)
    section.left_margin=Cm(2.2);section.right_margin=Cm(2.0)
    section.footer_distance=Cm(.8)
    for name,size in [('Normal',12),('Title',20),('Subtitle',13),('Heading 1',15),('Heading 2',12.5),('Caption',10)]:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
        fonts=style.element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            if 'theme' in key.lower():del fonts.attrib[key]
        fonts.set(qn('w:cs'),'Times New Roman')
        for border in list(style.element.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
        style.paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing=1.12
    doc.styles['Normal'].paragraph_format.widow_control=True
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].font.bold=True;doc.styles[name].paragraph_format.keep_with_next=True
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.title='Dự đoán xác suất kết quả Premier League mùa giải 2024 2025'
    doc.core_properties.subject='Khai phá và Phân tích dữ liệu'
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    md=[]
    def heading(text,new_page=True):
        if new_page:doc.add_page_break()
        doc.add_heading(text,level=1);md.extend(['',f'## {text}',''])
    def para(text):doc.add_paragraph(text);md.extend([text,''])
    def table(headers,rows,widths=None):
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        widths=widths or [16.8/len(headers)]*len(headers)
        for col,width in zip(t.columns,widths):col.width=Cm(width)
        for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            cells=t.add_row().cells
            for cell,text in zip(cells,values):cell.text=str(text)
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        t._tbl.tblPr.append(borders)
        for i,row in enumerate(t.rows):
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
            for cell,width in zip(row.cells,widths):
                cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E5EDF3' if i==0 else ('F6F8FA' if i%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shading)
                margins=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    v=OxmlElement('w:'+side);v.set(qn('w:w'),'70');v.set(qn('w:type'),'dxa');margins.append(v)
                cell._tc.get_or_add_tcPr().append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.0
                    for run in p.runs:run.font.size=Pt(10);run.bold=i==0;run.font.color.rgb=RGBColor(0,0,0)
        doc.add_paragraph().paragraph_format.space_after=Pt(2)
        md.append('| '+' | '.join(map(str,headers))+' |');md.append('| '+' | '.join(['---']*len(headers))+' |')
        md.extend('| '+' | '.join(map(str,row))+' |' for row in rows);md.append('')
    def figure(relative,caption,width=16.5):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        run=p.add_run();run.add_picture(str(ROOT/relative),width=Cm(width))
        inline=run._r.xpath('.//wp:docPr')
        if inline:inline[0].set('descr',caption)
        c=doc.add_paragraph(caption,style='Caption');c.alignment=WD_ALIGN_PARAGRAPH.CENTER
        md.extend([f'![{caption}](../../{relative})',''])
    def math_run(text):
        e=OxmlElement('m:r');v=OxmlElement('m:t');v.text=text;e.append(v);return e
    def subscript(base,items):
        e=OxmlElement('m:sSub');body=OxmlElement('m:e');body.append(math_run(base));sub=OxmlElement('m:sub')
        for item in items:sub.append(math_run(item) if isinstance(item,str) else item)
        e.extend([body,sub]);return e
    def summation(variable,upper,items):
        e=OxmlElement('m:nary');pr=OxmlElement('m:naryPr');char=OxmlElement('m:chr');char.set(qn('m:val'),'∑')
        loc=OxmlElement('m:limLoc');loc.set(qn('m:val'),'undOvr');pr.extend([char,loc])
        lower=OxmlElement('m:sub');lower.append(math_run(variable+'=1'))
        top=OxmlElement('m:sup');top.append(math_run(upper));body=OxmlElement('m:e')
        for item in items:body.append(item)
        e.extend([pr,lower,top,body]);return e
    def equation(label,kind):
        # Native editable Word math, including real fractions rather than raw LaTeX.
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        math=OxmlElement('m:oMath');math.append(math_run(label+' = '))
        frac=OxmlElement('m:f');num=OxmlElement('m:num');den=OxmlElement('m:den')
        if kind=='temperature':
            math.append(math_run('softmax('));num.append(math_run('log p'));den.append(math_run('T'))
        else:num.append(math_run('1'));den.append(math_run('N'))
        frac.append(num);frac.append(den)
        if kind=='logloss':math.append(math_run('−'))
        math.append(frac)
        if kind=='temperature':math.append(math_run(')'))
        elif kind=='logloss':math.append(summation('i','N',[math_run('ln '),subscript('p',['i,',subscript('y',['i'])])]))
        else:
            power=OxmlElement('m:sSup');body=OxmlElement('m:e')
            body.extend([math_run('('),subscript('p',['i,c']),math_run('−'),subscript('y',['i,c']),math_run(')')])
            exponent=OxmlElement('m:sup');exponent.append(math_run('2'));power.extend([body,exponent])
            math.append(summation('i','N',[summation('c','3',[power])]))
        p._p.append(math)
        md.append({'logloss':'Log Loss = −mean(log xác suất của nhãn đúng).','brier':'Brier = mean(tổng bình phương sai số ba xác suất so với vector one-hot).','temperature':'p_calibrated = softmax(log(p) / T).'}[kind]);md.append('')
    def metric_row(label,model,variant='raw'):
        m=metrics[model+'::'+variant]
        return [label,f"{m['accuracy']:.2%}",f"{m['macro_f1']:.4f}",f"{m['log_loss']:.5f}",f"{m['brier_sum']:.5f}"]
    # 1 opening
    title='Dự đoán xác suất kết quả Premier League mùa giải 2024 2025'
    doc.add_paragraph(title,style='Title');md.extend(['# '+title,''])
    doc.add_paragraph('Thống kê đội bóng và các đặc trưng lịch sử cầu thủ',style='Subtitle')
    para('Bài tập lớn môn Khai phá và Phân tích dữ liệu\nBáo cáo thực nghiệm hoàn thành ngày 24 tháng 10 năm 2026')
    doc.add_heading('Tóm tắt',level=1);md.extend(['## Tóm tắt',''])
    para('Nghiên cứu xây dựng hệ thống dự đoán xác suất ba kết quả chủ nhà thắng, hòa và đội khách thắng của từng trận Premier League. Dữ liệu miễn phí được kiểm định và chuẩn hóa trước khi tạo thống kê theo thời gian. Nhánh A sử dụng phong độ và Elo cấp đội; P bổ sung lịch sử sử dụng nhân sự; F bổ sung hiệu suất bàn thắng và kiến tạo của nhóm cầu thủ tham gia trận trước. Các nhánh cầu thủ được so sánh trên cùng tập trận để giảm sai lệch do khác giai đoạn dữ liệu.')
    para('Trên 380 trận test mùa 2024–2025, Elo A_full đã được chọn trước test làm tham chiếu đạt Accuracy 54,74%, Macro F1 0,41096 và Log Loss gốc 0,98141. Temperature scaling fit trên validation 2023–2024 làm Log Loss test tăng ở cả 14 estimator. Sáu so sánh chính về P/F chưa cho thấy cải thiện chắc chắn; mọi khoảng điều chỉnh cho nhiều so sánh tại block chính đều chứa 0. Kết quả nhấn mạnh giá trị của baseline mạnh và đánh giá xác suất ngoài mẫu.')
    para('Hệ thống là backtest hồi cứu theo ngày. Chưa xác minh announced XI và thời điểm nguồn công bố tại T−60, nên không suy diễn kết quả proxy thành kết luận về đội hình chính thức hoặc giá trị thị trường. Đóng góp trọng tâm là quy trình chống rò rỉ, audit có bằng chứng, ablation công bằng và lịch sử thí nghiệm truy vết được.')
    table(['Quy mô','Giá trị'],[['Dữ liệu đội','12 mùa gồm 2 mùa warm-up'],['Lịch sử cầu thủ chuẩn hóa','128.543 dòng development, 1.900 trận'],['Vòng test đã khóa','380 trận, 30 đầu ra xác suất'],['Kiểm thử trước test','81 unit tests và các kiểm tra thời gian trên dữ liệu thật']],[7.0,9.8])
    # 2 motivation
    heading('1 Vấn đề nghiên cứu và cơ sở phương pháp')
    para('Mục tiêu là ước lượng xác suất cho H, D, A thay vì chỉ chọn một nhãn. Hai trận đều được dự đoán H có thể có mức chắc chắn rất khác nhau; vì vậy Accuracy chưa đủ phản ánh chất lượng dự đoán. Nghiên cứu ưu tiên Log Loss và Brier, đồng thời báo chất lượng phân loại từng lớp.')
    para('Câu hỏi khả thi với dữ liệu đã thu thập là liệu lịch sử sử dụng và đóng góp cầu thủ có bổ sung tín hiệu cho thống kê đội hay không. Câu hỏi ban đầu về XI chính thức, tuổi và market value cần bằng chứng thời điểm công bố mà các bảng tải được chưa cung cấp. Nghiên cứu giữ phần đó thành giới hạn dữ liệu, không thay thế bằng thông tin sau trận rồi coi là đầu vào trước trận.')
    para('Logistic Regression đa lớp là đối chứng dễ kiểm toán và có regularization. Random Forest kiểm tra quan hệ phi tuyến mà không cần chuẩn hóa thang đo. XGBoost là thuật toán boosting trên cây được mô tả bởi Chen và Guestrin [4]. Quy mô vài nghìn trận không tạo lý do bắt buộc phải dùng Deep Learning; nghiên cứu giới hạn số cấu hình để kiểm soát chi phí và mức thử nghiệm.')
    para('Guo và cộng sự trình bày temperature scaling như phương pháp hậu xử lý một tham số [5]. Kết quả của họ không bảo đảm phương pháp luôn cải thiện trên dữ liệu bóng đá theo mùa. Trong dự án, phương pháp được fit trên mùa validation rồi đánh giá ngoài mẫu ở mùa tiếp theo. Pipeline chỉ fit các phép tiền xử lý trên training theo hướng dẫn kiểm soát data leakage của scikit-learn [3].')
    table(['Câu hỏi','Thực nghiệm trả lời'],[['Thống kê đội dự đoán được đến đâu','A_full và các baseline xác suất'],['Mức sử dụng cầu thủ có bổ sung tín hiệu','A_common so với A+P'],['Phong độ đóng góp có bổ sung thêm','A+P so với A+P+F'],['Calibration có tổng quát hóa sang mùa sau','Raw và temperature trên test'],['XI và market value chính thức có tác động không','Chưa đủ điều kiện dữ liệu để kết luận']],[7.0,9.8])
    para('Không tối ưu Accuracy bằng cách thay ngưỡng lớp hòa sau khi xem test. Các lỗi dự đoán hòa, kết quả không cải thiện và sự suy giảm sau calibration đều được công bố.')
    # 3 sources
    heading('2 Dữ liệu và kiểm định chất lượng')
    para('Football-Data.co.uk cung cấp CSV trận EPL cho 2013–2014 đến 2024–2025 trong bộ dữ liệu tải thực tế [1]. Mỗi mùa có 380 trận sau chuẩn hóa; hai mùa đầu phục vụ warm-up. Nguồn cầu thủ là archive FPL của vaastav, được tải và lưu checksum từng file [2]. Giá FPL không đồng nhất với market value. Các nguồn FBref, Transfermarkt/Kaggle và football-data.org được nghiên cứu ở giai đoạn khả thi nhưng không được dùng để giả định có XI hoặc valuation lịch sử hợp lệ.')
    table(['Mùa cầu thủ','Dòng raw','Tình trạng chuẩn hóa'],[['2016–2017','23.679','Thiếu tệp fixtures và teams ở đường dẫn kiểm tra'],['2017–2018','22.467','Thiếu tệp fixtures và teams ở đường dẫn kiểm tra'],['2018–2019','21.790','Thiếu tệp teams ở đường dẫn kiểm tra'],['2019–2020','22.560','22.501 dòng sau loại placeholder, ghép 380 trận'],['2020–2021','24.365','Ghép đủ 380 trận'],['2021–2022','25.447','Ghép đủ 380 trận'],['2022–2023','26.505','Ghép đủ 380 trận, advanced stats có khoảng thiếu'],['2023–2024','29.725','Ghép đủ 380 trận validation']],[3.0,2.3,11.5])
    para('Mùa 2019–2020 có 59 khóa player–fixture lặp tại trận Man City–Arsenal bị hoãn. Bản ngày 11/03/2020 có minutes=0 và tỷ số trống; từng cầu thủ có bản thay thế ở trận ngày 17/06/2020. Chỉ loại placeholder khi thỏa đầy đủ quy tắc đã ghi và đối chiếu thông báo Premier League [7]. Raw vẫn giữ nguyên; không dùng thao tác giữ dòng cuối tùy tiện.')
    para('Trong 136 trận đầu mùa 2022–2023, starts, xG và xA của toàn trận đều bằng 0; chỉ 244 trận có tổng 22 recorded starts. Bộ ba toàn 0 bị đánh dấu không dùng được ở bảng chuẩn hóa, còn số 0 hợp lệ của một cá nhân được giữ. P/F không dùng các cột này. Test 2024–2025 loại thêm 322 dòng Assistant Manager trước khi tạo feature.')
    para('Đội tại trận được lấy từ fixture và was_home, không từ đội cuối mùa của cầu thủ. ID cầu thủ có season scope. Tên, mùa, đội và ngày trận được ghép có kiểm tra một-một hoặc nhiều-một rõ ràng. Khác giờ trong cùng ngày được giữ cả hai bản và gắn cờ; không tự coi lịch archive là lịch đã biết tại thời điểm dự đoán.')
    # 4 temporal data
    heading('3 Dataset và kiểm soát rò rỉ theo thời gian')
    table(['Nhánh','Training','Validation','Testing'],[['A_full','2015–2023, 3.040 trận','2023–2024, 380','2024–2025, 380'],['A_common','2019–2023, 1.520 trận','Cùng 380 trận','Cùng 380 trận'],['A+P','Cùng 1.520 trận','Cùng 380 trận','Cùng 380 trận'],['A+P+F','Cùng 1.520 trận','Cùng 380 trận','Cùng 380 trận']],[3.0,5.2,4.3,4.3])
    para('Luồng xử lý gồm raw bất biến, bảng chuẩn hóa có ID và thời gian, rồi feature snapshots. Match outcomes tách khỏi ma trận đầu vào. Mỗi snapshot có lineage lưu trận nguồn và availability giả định lớn nhất. Dự đoán được tạo trước khi cập nhật trạng thái từ kết quả cả ngày; mọi trận trong cùng ngày dùng thông tin từ các ngày trước.')
    para('Nếu chỉ có ngày thi đấu, không tự gán giờ 15:00. Bản tái dựng dùng 00:00 Europe/London của ngày hiện tại làm mốc trạng thái và giả định thông tin trận trước có vào nửa đêm kế tiếp. Quy tắc này bảo thủ với các trận cùng ngày nhưng vẫn chưa chứng minh publication time lịch sử. Múi giờ địa phương xử lý DST, không cộng cố định 24 giờ theo UTC.')
    para('Elo và form cấp đội được cập nhật tuần tự. Phong độ có thể đi qua hai mùa liên tiếp của cùng đội; đội mới lên hạng được khởi tạo theo quy tắc cố định. P/F reset đầu mùa vì chưa nối identity cầu thủ xuyên mùa. Không backfill dữ liệu tương lai, không dùng thứ hạng cuối mùa hoặc thống kê của trận đang dự đoán.')
    para('Days since last PL match chỉ đo lịch EPL, không được gọi là thời gian nghỉ ở mọi đấu trường. Sporting points là 3W+D, chưa bao gồm lịch sử quyết định trừ điểm. Những giới hạn này được thể hiện bằng tên feature và từ điển dữ liệu, thay vì trình bày như trạng thái chính thức đầy đủ.')
    table(['Nhóm kiểm tra','Điều kiện phải đúng'],[['Trận hiện tại','Đổi outcome hoặc minutes/goals/assists không đổi feature của chính trận đó'],['Tương lai','Thay dữ liệu các trận sau không làm đổi feature trước đó'],['Prefix và thứ tự','Tính trên prefix khớp full history; đảo input không đổi kết quả'],['Join và lineage','ID duy nhất, đúng phía đội, không vượt cutoff giả định'],['Tiền xử lý','Imputer/scaler chỉ fit trên training của từng fold']],[4.2,12.6])
    # 5 features
    heading('4 Biểu diễn đội bóng và lịch sử cầu thủ')
    para('A có 40 cột số gồm Elo home/away và chênh lệch, phong độ 5/10 trận, hiệu suất sân nhà/sân khách, sporting points/PPG/GD trước trận, số trận hỗ trợ, ngày từ trận EPL gần nhất, cờ thiếu lịch sử/mới lên hạng và H2H. Không có odds, nhãn hiện tại hoặc thống kê phát sinh trong trận.')
    table(['Đại lượng P mỗi phía','Ý nghĩa trong tối đa 5 trận trước'],[['Số trận lịch sử','0–5, hỗ trợ diễn giải missing'],['Số người có phút dương','Số player-season ID từng tham gia'],['HHI phút thi đấu','Mức tập trung tổng phút vào ít người'],['Tỷ trọng phút top 11','Tổng tỷ trọng phút của 11 người được dùng nhiều nhất'],['Tỷ trọng phút người quay lại','Phút trận trước thuộc người đã tham gia sớm hơn trong cửa sổ'],['Overlap phút trung bình','Tổng min của tỷ trọng phút qua từng cặp trận liền nhau']],[6.0,10.8])
    para('P có 12 cột vì tính sáu đại lượng cho hai phía. Không có lịch sử thì số trận bằng 0 và các đại lượng khác missing; continuity cần ít nhất hai trận. Nhóm top 11 là mô tả lịch sử mức sử dụng, không phải starting XI hiện tại. Dòng không có phút dương không tạo người tham gia trong cửa sổ.')
    para('F thêm bốn đại lượng mỗi phía. Rate cá nhân lấy tổng bàn hoặc kiến tạo trong cửa sổ nhân 90, rồi chia tổng phút cộng 180. Rate nhóm là tổng có trọng số, trong đó trọng số là tỷ trọng phút của người đó ở trận trước gần nhất. 180 phút là prior cố định về rate 0, ghi trước khi đo F; mục đích là tránh per90 rất lớn từ mẫu ít phút.')
    para('Hai đại lượng F còn lại là HHI đóng góp G+A và phần đóng góp trong cửa sổ thuộc nhóm người tham gia trận trước gần nhất. Tổng đóng góp bằng 0 khiến các tỷ trọng missing; đây không phải số 0 quan sát được. Không có lịch sử thì cả bốn missing. Mỗi bên được xử lý cùng công thức, không dùng current roster để chọn người.')
    para('Assists sử dụng định nghĩa FPL. P/F chỉ dùng lịch sử của chính đội, không tự mang phong độ xuyên CLB hoặc qua mùa. Người chuyển đi có thể vẫn nằm trong lịch sử; thiếu roster đúng thời điểm làm proxy có sai số. Các cột xG/xA, starts, tuổi, market value và chấn thương được loại khỏi các phiên bản đã thực nghiệm.')
    # 6 models/eval
    heading('5 Mô hình đánh giá và calibration')
    table(['Thuật toán','Bốn ứng viên trong mỗi nhánh'],[['Logistic Regression','C ∈ {0,01; 0,1; 1; 10}, L2, multinomial'],['Random Forest','max_depth ∈ {4;8} × min_samples_leaf ∈ {10;30}, 300 cây'],['XGBoost','max_depth ∈ {2;3} × n_estimators ∈ {100;250}, learning_rate 0,03']],[4.1,12.7])
    para('Mỗi nhánh chọn tham số riêng bằng Log Loss trung bình ba fold expanding-window trong training. Với common: train 380, 760 và 1.140 trận để đánh giá từng mùa 2020–2021, 2021–2022 và 2022–2023, mỗi fold 380 trận. Median imputation và missing indicators được fit riêng từng fold; Logistic thêm StandardScaler. Không class weighting, không early stopping bằng validation cuối, không đổi ngưỡng theo test.')
    para('Baseline xác suất gồm tần suất H/D/A từ training với Laplace smoothing và Logistic ba lớp trên riêng elo_diff. Expected score của Elo không được coi trực tiếp là xác suất thắng do còn lớp hòa. Luôn chọn H chỉ có metric phân loại, không gán xác suất 0 để đánh giá Log Loss.')
    equation('Log Loss','logloss');equation('Brier','brier')
    para('N là số trận; p là xác suất theo thứ tự H,D,A. Trong Log Loss, y là nhãn đúng; trong Brier, y theo hai chỉ số là biến chỉ báo one-hot của lớp. Brier dùng tổng ba lớp rồi lấy trung bình, miền [0,2]. Accuracy, Macro F1, precision/recall từng lớp, confusion matrix và reliability bổ sung việc đánh giá xác suất. Nhãn được lấy bằng argmax.')
    equation('p hiệu chỉnh','temperature')
    para('Fit một T cho mỗi estimator trên 380 trận validation 2023–2024, với T trong [0,5;5] và có phương án T=1. Điểm validation sau fit là diagnostic in-sample. Base estimators không refit lên validation. Temperature giữ argmax nên không tự sửa recall hòa [6]. Hai baseline tần suất giữ raw.')
    para('So sánh ghép cặp dùng moving-block bootstrap theo tuần lịch, 2.000 lần, seed 42. Block chính hai tuần; một/bốn tuần là độ nhạy. Sáu contrast temperature là phân tích chính: P trừ A và APF trừ AP cho ba thuật toán. Báo CI 95% cùng khoảng Bonferroni-style 99,1667%; độ phủ chỉ xấp xỉ và số mẫu đuôi hữu hạn.')
    # 7 validation
    heading('6 Kết quả development trên validation')
    para('Ở A_full, Elo đạt Log Loss 0,92848; Random Forest 0,92986; XGBoost 0,93882; Logistic 0,94166. Elo A_full được chọn làm tham chiếu thực dụng trước khi mở test. Mọi so sánh P/F bên dưới dùng cohort common train 1.520 trận; không quy chênh lệch với A_full train 3.040 trận thành tác dụng của feature.')
    previous=json.loads((ROOT/'reports/APF_dev_v1/results.json').read_text())['validation_metrics']
    rows=[]
    for family in ['logistic','random_forest','xgboost']:
        for arm,display in [('A_common','A'),('A_plus_P','A+P'),('A_plus_P_plus_F','A+P+F')]:
            m=previous[f'{family}__{arm}'];rows.append([FAMILY[family],display,f"{m['accuracy']:.2%}",f"{m['macro_f1']:.4f}",f"{m['log_loss']:.5f}"])
    table(['Thuật toán','Nhánh','Accuracy','Macro F1','Log Loss raw'],rows,[4.4,2.3,3.2,3.2,3.7])
    para('P chưa cải thiện Log Loss rõ ràng. XGBoost A+P tăng Accuracy lên 58,16% nhưng mức giảm Log Loss nhỏ; cả chín CI bootstrap theo block 1/2/4 tuần đều chứa 0. Thêm F làm Logistic kém hơn trên điểm ước lượng, trong khi RF/XGBoost giảm Log Loss rất ít so với A+P. Đây là lý do cần giữ metric xác suất làm tiêu chí chính.')
    para('Feature importance theo nhóm được thực hiện trên validation bằng hoán vị các cột cùng nhóm trong từng tháng. Elo là nhóm có tác động lớn nhất trong mô hình A đã fit. Tác động P/F nhỏ hoặc không ổn định; các đặc trưng tương quan có thể thay thế nhau và hoán vị có thể tạo tổ hợp ít thực tế. Không diễn giải các giá trị này như quan hệ nhân quả.')
    para('Development diễn ra theo từng phiên bản sau khi đã xem validation ở các bước trước. Mỗi lượt giữ cấu hình ghi trước fit, nhưng không gọi validation là tập chưa từng quan sát. Sau phiên bản F đã đăng ký, dự án dừng mở rộng cửa sổ, shrinkage, lưới hoặc thuật toán để tìm điểm tốt hơn; mọi kết quả không cải thiện được giữ lại.')
    # 8 final results
    heading('7 Kết quả test cuối mùa 2024 2025')
    para('Trước scoring, danh sách 14 estimator, hai baseline, 14 nhiệt độ, feature và mã được khóa bằng SHA-256. Commit khóa là 7773397; kết quả được lưu sau đó tại b36fb63. Runner lưu đủ xác suất trước khi ghép nhãn để tính metric. Không refit estimator/calibrator trong test; trạng thái lịch sử được cập nhật từ các ngày test trước theo quy tắc đã chốt.')
    table(['Tham chiếu Elo A_full','Raw','Temperature'],[['Accuracy','54,74%','54,74%'],['Macro F1','0,41096','0,41096'],['Log Loss','0,98141','0,98606'],['Brier','0,58654','0,58993']],[7.2,4.8,4.8])
    rows=[]
    for entry in data['registry']:
        raw=metrics[entry['id']+'::raw'];cal=metrics.get(entry['id']+'::temperature')
        arm={'A_full':'Full','A_common':'A','A_plus_P':'A+P','A_plus_P_plus_F':'APF'}[entry['arm']]
        scope='Full' if entry['cohort']=='full' else 'Common'
        rows.append([FAMILY[entry['family']].replace(' Regression',''),scope+'/'+arm,f"{raw['accuracy']:.2%}",f"{raw['log_loss']:.5f}",f"{cal['log_loss']:.5f}" if cal else '—'])
    table(['Thuật toán','Cohort và nhánh','Accuracy','LL raw','LL temp'],rows,[4.6,4.0,2.6,2.8,2.8])
    para('Test có H/D/A = 155/93/132. Luôn chọn chủ nhà đạt 40,79% Accuracy. Cả 14 mô hình có Log Loss và Brier tăng sau temperature. Không chọn lại raw/calibrated sau khi thấy test; bảng giữ cả hai theo protocol. Chi tiết Brier, F1 và confusion matrices mọi mô hình nằm trong báo cáo kết quả đầy đủ.')
    # 9 contrast
    heading('8 Đóng góp bổ sung của P và F trên test')
    para('So sánh chính sử dụng temperature đã fit trước test. Delta là Log Loss mô hình thêm feature trừ đối chứng cùng thuật toán và cohort; âm có lợi cho feature mới. Cả sáu khoảng chính, kể cả CI 95%, đều chứa 0. Không coi điều này là bằng chứng hai mô hình tương đương hoặc mọi dữ liệu cầu thủ đều không hữu ích.')
    primary=[u for u in r['paired_feature_comparisons'] if u['variant']=='temperature' and u['block_weeks']==2]
    rows=[]
    for u in primary:
        kind='P − A' if u['model_arm']=='A_plus_P' else 'APF − AP'
        rows.append([FAMILY[u['family']].replace(' Regression',''),kind,f"{u['delta_log_loss']:+.5f}",f"[{u['ci95_low']:.5f}; {u['ci95_high']:.5f}]",f"[{u['family_ci_low']:.5f}; {u['family_ci_high']:.5f}]"])
    table(['Thuật toán','Contrast','Delta','CI 95%','CI 99,1667%'],rows,[3.1,2.0,2.1,4.8,4.8])
    figure('reports/final_test_v1/feature_comparison.png','Hình 1 So sánh feature và độ bất định trên 380 trận test')
    para('P làm Log Loss tăng ở cả ba thuật toán trên ước lượng điểm. F làm giảm Log Loss so với A+P, nhưng mức giảm khoảng 0,0037–0,0042 chưa tách rõ khỏi nhiễu của một mùa. Kết quả raw và độ nhạy block một/bốn tuần đều được lưu, không chọn block thuận lợi để thay kết luận chính.')
    para('Các khoảng giữ cố định mô hình, chưa bao gồm bất định do huấn luyện lại hoặc biến động giữa nhiều mùa. Hiệu chỉnh Bonferroni-style không làm các giả định bootstrap trở thành đúng tuyệt đối; 2.000 resample cũng tạo giới hạn khi ước lượng đuôi 99,1667%.')
    # 10 interpretation
    heading('9 Thảo luận giới hạn và kết luận')
    para('Elo A_full không chọn nhãn hòa lần nào, recall D bằng 0 dù test có 93 trận hòa. Temperature giữ nguyên nhãn nên không giải quyết hạn chế này. Xác suất D vẫn được tính đầy đủ vào Log Loss/Brier; việc không chọn D không đồng nghĩa p(D)=0. Báo cáo không che điểm yếu bằng Accuracy tổng.')
    figure('reports/final_test_v1/reference_reliability.png','Hình 2 Reliability theo lớp của tham chiếu Elo A_full',16.2)
    para('Tất cả T dưới 1 làm dự đoán sắc hơn theo validation, nhưng ngoài mẫu lại tăng loss. Tỷ lệ chủ nhà thắng giảm từ 46,05% ở validation xuống 40,79% ở test, hòa tăng từ 21,58% lên 24,47%. Đây là khác biệt mô tả giữa hai mùa, chưa chứng minh nguyên nhân calibration thất bại. Quan hệ giữa đặc trưng và nhãn cũng có thể thay đổi; một mùa calibration nhỏ không bảo đảm tổng quát hóa.')
    para('Các giới hạn dữ liệu gồm chưa xác minh publication time, thiếu announced XI và market value đúng thời điểm, không nối player identity xuyên mùa, roster hiện tại không biết đầy đủ, assists theo FPL và rest days chỉ theo EPL. Dữ liệu archive có thể đã được sửa về sau. Vì vậy kết quả là backtest hồi cứu có giả định, không phải chứng minh hệ thống live tại T−60 với feed đầy đủ.')
    para('Kết luận được hỗ trợ là: trong cohort và các phiên bản đã định nghĩa, proxy nhân sự/đóng góp cầu thủ chưa tạo cải thiện Log Loss chắc chắn; baseline Elo là đối chứng mạnh và temperature fit mùa trước làm chất lượng xác suất test giảm trong thí nghiệm này. Không suy rộng thành “cầu thủ không quan trọng” hoặc “calibration luôn có hại”. Hướng mở rộng dữ liệu tương lai là XI có timestamp, roster/identity và nguồn performance nhất quán; phải dùng một giai đoạn đánh giá mới.')
    # 11 reproducibility/ref
    heading('10 Tái lập sản phẩm và tài liệu tham khảo')
    para('Bộ bàn giao có báo cáo Word và Markdown, notebook đã thực thi, demo HTML offline và các artifact kết quả/khóa thí nghiệm. Notebook tự tính lại metric từ xác suất đã lưu; demo chỉ đổi góc xem trên archive. Cả hai không huấn luyện hoặc chọn lại mô hình. Full reproduction từ raw cần đúng snapshot, phiên bản thư viện và model artifacts, không thể đồng nhất với việc đọc lại CSV kết quả.')
    para('Đã có 81 unit tests trước vòng test và các kiểm tra prefix, current/future perturbation, thứ tự input, DST và lineage trên dữ liệu thật. Nạp lại artifact cuối tái tạo đủ 11.400 bộ xác suất cùng mọi metric/confusion matrix, sai số xác suất tối đa khoảng 3,33×10⁻¹⁶. Notebook có 10 ô mã chạy thành công; demo được kiểm tra desktop/mobile, bộ lọc/empty state, ẩn kết quả thật và không gọi mạng.')
    para('Để trình bày: mở demo/index.html, chọn trận/mô hình, đọc xác suất raw và temperature, rồi bật kết quả thực tế. Khi giải thích, nhấn mạnh khác biệt giữa dự đoán xác suất và kết quả được bảo đảm, giữa proxy nhóm trận trước và đội hình chính thức. Không sửa mô hình từ lỗi test để tái dùng mùa này như một tập chưa từng quan sát.')
    refs=[
      '[1] Football-Data.co.uk. England historical results CSV. https://www.football-data.co.uk/englandm.php',
      '[2] Vaastav. Fantasy Premier League historical archive và LICENSE. https://github.com/vaastav/Fantasy-Premier-League',
      '[3] scikit-learn. Common pitfalls and recommended practices, Data leakage. https://scikit-learn.org/stable/common_pitfalls.html',
      '[4] Chen, T. và Guestrin, C. (2016). XGBoost A Scalable Tree Boosting System. https://arxiv.org/abs/1603.02754',
      '[5] Guo, C., Pleiss, G., Sun, Y. và Weinberger, K. Q. (2017). On Calibration of Modern Neural Networks. PMLR 70, 1321–1330. https://proceedings.mlr.press/v70/guo17a.html',
      '[6] scikit-learn. Probability calibration, Temperature scaling. https://scikit-learn.org/stable/modules/calibration.html',
      '[7] Premier League (2020). Man City v Arsenal postponed và initial fixtures for season restart. https://www.premierleague.com/en/news/1643922 ; https://www.premierleague.com/en/news/1676717',
      '[8] Artifact nội bộ dự án. final_evaluation_config.json, final_model_registry.json, final_lock.json; reports/final_test_v1/results.json và artifact_verification.json. Khóa trước scoring: commit 7773397.']
    for text in refs:
        p=doc.add_paragraph(text);p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(5)
        for run in p.runs:run.font.size=Pt(10)
        md.extend([text,''])
    para('Thời điểm kiểm chứng nguồn và thực nghiệm: tháng 10 năm 2026. Các con số là kết quả thực nghiệm đã lưu, không phải dữ liệu minh họa. Raw và model artifacts giữ cục bộ; giấy phép mã nguồn không mặc nhiên tái cấp phép dữ liệu bên thứ ba.')
    OUT.mkdir(exist_ok=True)
    doc.save(OUT/'academic_report.vi.docx')
    (OUT/'academic_report.vi.md').write_text('\n'.join(md),encoding='utf-8')
    print('Created academic_report.vi.docx and its Markdown source.')


if __name__=='__main__':main()
